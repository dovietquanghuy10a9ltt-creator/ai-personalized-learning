from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel
import unicodedata 
import json
import io
import random
import copy
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from db.database import get_db
from db import models
from groq import Groq
import os
from rag.vector_store import get_vector_store

router = APIRouter()

class ExamRequest(BaseModel):
    class_id: int
    subject: str
    exam_type: str # "trắc nghiệm" hoặc "tự luận"
    num_questions: int
    num_versions: int # Số lượng mã đề
    level: str 

def create_exam_header(doc, subject, exam_type, exam_code):
    """Hàm tạo Header chuẩn form Trường Đại học Xây Dựng Hà Nội"""
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    
    cell_left = table.cell(0, 0)
    p_left1 = cell_left.paragraphs[0]
    p_left1.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run_left1 = p_left1.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC XÂY DỰNG HÀ NỘI")
    run_left1.font.size = Pt(11)
    
    p_left2 = cell_left.add_paragraph()
    p_left2.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run_left2 = p_left2.add_run("KHOA CÔNG NGHỆ THÔNG TIN")
    run_left2.font.size = Pt(12)
    run_left2.bold = True
    
    p_left3 = cell_left.add_paragraph()
    p_left3.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p_left3.add_run("-----------------------").bold = True

    cell_right = table.cell(0, 1)
    p_right1 = cell_right.paragraphs[0]
    p_right1.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run_right1 = p_right1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    run_right1.font.size = Pt(11)
    run_right1.bold = True
    
    p_right2 = cell_right.add_paragraph()
    p_right2.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run_right2 = p_right2.add_run("Độc lập - Tự do - Hạnh phúc")
    run_right2.font.size = Pt(12)
    run_right2.bold = True
    
    p_right3 = cell_right.add_paragraph()
    p_right3.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p_right3.add_run("-----------------------").bold = True

    doc.add_paragraph() 
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run_title = p_title.add_run("ĐỀ THI KẾT THÚC HỌC PHẦN")
    run_title.font.size = Pt(16)
    run_title.bold = True

    p_info = doc.add_paragraph()
    p_info.add_run(f"Môn thi: ").bold = True
    p_info.add_run(f"{subject}\n")
    p_info.add_run(f"Hình thức thi: ").bold = True
    p_info.add_run(f"{exam_type.title()}\n")
    p_info.add_run(f"Mã đề thi: ").bold = True
    p_info.add_run(f"{exam_code}\n")
    p_info.add_run("Họ và tên sinh viên: ..................................................... MSSV: ............................... Lớp: ..................")
    
    doc.add_paragraph("--------------------------------------------------------------------------------------------------------------------------")

def remove_accents(input_str):
    s = input_str.replace('đ', 'd').replace('Đ', 'D')
    nfkd_form = unicodedata.normalize('NFKD', s)
    return u"".join([c for c in nfkd_form if not unicodedata.combining(c)])

@router.post("/generate-word")
def generate_exam_word(req: ExamRequest, db: Session = Depends(get_db)):
    # 1. RAG - Lấy tài liệu gốc cực lớn
    docs = db.query(models.Document).filter(
        models.Document.class_id == req.class_id,
        models.Document.subject == req.subject
    ).all()
    allowed_filenames = [doc.filename for doc in docs]
    
    context_summary = "Sử dụng kiến thức chuyên ngành chuẩn xác."
    if allowed_filenames:
        try:
            vector_store = get_vector_store()
            # Tăng lượng tài liệu nạp vào để AI không bị bí ý tưởng
            search_results = vector_store.similarity_search(
                f"Toàn bộ tài liệu, slide bài giảng, code mẫu môn {req.subject}", k=40, filter={"source": {"$in": allowed_filenames}}
            )
            context_summary = "\n".join([d.page_content for d in search_results])[:25000]
        except Exception as e:
            print("Lỗi RAG:", e)

    # 2. PROMPT CHO AI TẠO BỘ CÂU HỎI (ĐỒNG BỘ VỚI ASSESSMENT AGENT)
    api_key = os.getenv("GROQ_KEY_ADAPTIVE") or os.getenv("GROQ_KEY_ASSESSMENT")
    client = Groq(api_key=api_key)
    
    # Ép AI sinh dư thêm vài câu để buffer lọc lỗi
    ask_count = req.num_questions + 3 

    # --- KHỐI LỆNH ĐỘ KHÓ ---
    level_instruction = ""
    if req.level.lower() == "beginner":
        level_instruction = "[MỨC ĐỘ DỄ - BEGINNER]: Tập trung vào nhận biết cú pháp, lý thuyết cơ bản, đọc hiểu."
    elif req.level.lower() == "intermediate":
        level_instruction = "[MỨC ĐỘ TRUNG BÌNH]: Yêu cầu phân tích luồng chạy, tìm lỗi sai logic, vận dụng."
    else:
        level_instruction = "[MỨC ĐỘ KHÓ - ADVANCED]: Tập trung vào thiết kế hệ thống, tối ưu hóa, bẫy phức tạp."

    if req.exam_type == "trắc nghiệm":
        prompt = f"""
        BẠN LÀ HỘI ĐỒNG RA ĐỀ THI ĐẠI HỌC. Nhiệm vụ: Soạn ĐỀ THI XUẤT BẢN THỰC TẾ môn "{req.subject}".
        {level_instruction}
        
        [TÀI LIỆU BÀI GIẢNG (BẮT BUỘC PHẢI DỰA VÀO 100%)]:
        {context_summary}

        [🔴 LỆNH CẤM TỬ HÌNH - VI PHẠM SẼ BỊ HỦY DIỆT]:
        1. CẤM BỊA ĐẶT VÍ DỤ BÊN NGOÀI: Tài liệu có gì thì hỏi cái đó. TUYỆT ĐỐI KHÔNG tự động dùng các ví dụ kinh điển trên mạng (như Animal, Dog, Cat) nếu trong tài liệu KHÔNG CÓ.
        2. CẤM LỖI ĐÁP ÁN TRỐNG: Mảng "options" KHÔNG ĐƯỢC chỉ chứa ["A", "B", "C", "D"]. Nội dung phải là câu trả lời thật sự (Ví dụ: "Biến cục bộ", "Tham số truyền vào").
        3. CHỐNG LẶP LẠI (ANTI-LOOP): TẤT CẢ câu hỏi phải khai thác vấn đề KHÁC NHAU. CẤM lặp lại cấu trúc câu hỏi ("Phân tích khái niệm...").
        4. VĂN PHONG ĐA DẠNG: Hãy dùng: "Xét đoạn mã...", "Trong mô hình...", "Khi hệ thống...".

        [KPI ĐAN XEN TƯ DUY]:
        1. THỰC HÀNH / BÀI TẬP: Đưa ĐOẠN MÃ CODE / THÔNG SỐ (dùng \\n). Hỏi output, điền mã khuyết.
        2. TÌNH HUỐNG: Bối cảnh dự án, hệ thống phần mềm doanh nghiệp.
        3. LÝ THUYẾT NÂNG CAO: So sánh bản chất, tìm mệnh đề SAI.

        [CẤU TRÚC JSON BẮT BUỘC (PHẢI TẠO ĐÚNG {ask_count} CÂU)]:
        {{
            "questions": [
                {{
                    "q": "Nội dung câu hỏi logic, có bối cảnh hoặc đoạn code rõ ràng (Dùng \\n)...",
                    "options": [
                        "Nội dung đáp án 1 thật sự (Không ghi chữ A)", 
                        "Nội dung đáp án 2 thật sự", 
                        "Nội dung đáp án 3 thật sự", 
                        "Nội dung đáp án 4 thật sự"
                    ],
                    "ans": "A",
                    "exp": "Giải thích ngắn gọn."
                }}
            ]
        }}
        """
    else:
        prompt = f"""Soạn {req.num_questions} câu hỏi TỰ LUẬN môn {req.subject}. 
        {level_instruction}
        Tài liệu: {context_summary}
        Đầu ra JSON: {{"questions": [{{"q": "Câu hỏi tự luận tình huống/bài tập sâu sắc?", "ans": "Đáp án/Bareme", "exp": "Gợi ý"}}]}}"""

    # 3. GỌI AI VỚI MODEL LỚN
    try:
        chat_completion = client.chat.completions.create(
            messages=[{"role": "system", "content": f"Output valid JSON only. Generate exactly {ask_count if req.exam_type == 'trắc nghiệm' else req.num_questions} questions."}, {"role": "user", "content": prompt}],
            model="llama-3.3-70b-versatile", 
            temperature=0.6, 
            max_tokens=8000, 
            response_format={"type": "json_object"}
        )
        
        raw_content = chat_completion.choices[0].message.content
        clean_content = re.sub(r'```json\s*|\s*```', '', raw_content, flags=re.IGNORECASE).strip()
        ai_data = json.loads(clean_content)
        
        raw_questions = ai_data.get("questions", [])
        
        # --- BỘ LỌC PYTHON LÀM SẠCH ĐÁP ÁN ---
        base_questions = []
        if req.exam_type == "trắc nghiệm":
            for q in raw_questions:
                if len(base_questions) >= req.num_questions:
                    break # Đủ số lượng thì dừng
                
                raw_options = q.get('options', [])
                if len(raw_options) < 4: continue
                
                is_garbage = False
                valid_options = []
                for opt in raw_options:
                    clean_opt = re.sub(r'^([A-D][\.\:\-\)])\s*', '', str(opt)).strip()
                    if len(clean_opt) <= 1 or clean_opt.upper() in ["A", "B", "C", "D"]:
                        is_garbage = True
                        break
                    valid_options.append(clean_opt)
                
                if is_garbage: continue
                
                # Cập nhật lại options sạch sẽ
                q['options'] = valid_options
                base_questions.append(q)
        else:
            base_questions = raw_questions[:req.num_questions]
            
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi AI sinh đề: {str(e)}")

    doc = Document()
    exam_versions = []

    # =========================================================
    # 4. THUẬT TOÁN TRỘN ĐỀ VÀ TẠO FILE WORD
    # =========================================================
    for v in range(req.num_versions):
        exam_code = str(random.randint(101, 999))
        
        shuffled_qs = copy.deepcopy(base_questions)
        random.shuffle(shuffled_qs) 
        
        exam_versions.append({"code": exam_code, "questions": shuffled_qs})
        
        create_exam_header(doc, req.subject, req.exam_type, exam_code)

        for i, q in enumerate(shuffled_qs):
            p = doc.add_paragraph()
            p.add_run(f'Câu {i+1}: ').bold = True
            p.add_run(q.get('q', ''))
            
            if req.exam_type == "trắc nghiệm":
                options = q.get('options', [])
                correct_opt_idx = ord(q.get('ans', 'A').upper()) - 65
                if 0 <= correct_opt_idx < len(options):
                    correct_opt_text = options[correct_opt_idx]
                else:
                    correct_opt_text = options[0]
                
                random.shuffle(options) 
                
                labels = ['A', 'B', 'C', 'D']
                new_correct_label = "A"
                for idx, opt in enumerate(options):
                    if opt == correct_opt_text:
                        new_correct_label = labels[idx]
                        
                    doc.add_paragraph(f"{labels[idx]}. {opt}", style='List Bullet')
                
                q['new_ans'] = new_correct_label
            else:
                for _ in range(4): doc.add_paragraph("................................................................................................................................")
            doc.add_paragraph()
            
        doc.add_page_break() 

    # =========================================================
    # 5. VẼ ĐÁP ÁN CHO GIÁO VIÊN
    # =========================================================
    ans_heading = doc.add_heading('HƯỚNG DẪN CHẤM & ĐÁP ÁN CHI TIẾT', 1)
    ans_heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph()

    for version in exam_versions:
        doc.add_paragraph(f"MÃ ĐỀ: {version['code']}").bold = True
        
        if req.exam_type == "trắc nghiệm":
            table = doc.add_table(rows=1, cols=6)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            for i in range(6): hdr_cells[i].text = 'Câu - Đáp án'
            
            row_cells = table.add_row().cells
            col_idx = 0
            for i, q in enumerate(version['questions']):
                if col_idx > 5:
                    row_cells = table.add_row().cells
                    col_idx = 0
                row_cells[col_idx].text = f"Câu {i+1}: {q.get('new_ans', q.get('ans'))}"
                col_idx += 1
            doc.add_paragraph()
        else:
            for i, q in enumerate(version['questions']):
                p = doc.add_paragraph()
                p.add_run(f'Câu {i+1}: ').bold = True
                doc.add_paragraph(q.get('ans', ''))
                doc.add_paragraph(f"Gợi ý chấm: {q.get('exp', '')}").italic = True
        doc.add_paragraph("----------------------------------------------------------------")

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    clean_subject = remove_accents(req.subject).replace(' ', '')
    clean_exam_type = remove_accents(req.exam_type).replace(' ', '')
    filename = f"DeThi_{clean_subject}_{clean_exam_type}.docx"
    
    return StreamingResponse(
        file_stream, 
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )